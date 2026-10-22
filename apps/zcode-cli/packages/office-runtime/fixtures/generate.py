"""Regenerate the sample documents used by tests.

Run from this directory:
  uv run --with python-docx==1.1.2 --with openpyxl==3.1.5 --with python-pptx==1.0.2 --with reportlab==4.2.5 generate.py
"""

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


def make_docx(path):
    doc = Document()
    doc.core_properties.title = "Quarterly Report"
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("This report summarises revenue for the third quarter.")
    doc.add_heading("Revenue", level=2)
    p = doc.add_paragraph("Revenue grew by ")
    p.add_run("12 percent").bold = True
    p.add_run(" compared with the previous quarter.")
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Region", "Revenue"), ("North", "120"), ("South", "95")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    doc.add_heading("Outlook", level=2)
    doc.add_paragraph("We expect steady growth. 预计下季度保持稳定增长。")
    for i in range(1, 4):
        doc.add_paragraph(f"Risk item {i}: supply chain delays.", style="List Bullet")
    q = doc.add_paragraph("See ")
    q.add_run("https://example.com/report").font.size = Pt(10)
    doc.add_page_break()
    doc.add_heading("Appendix", level=2)
    doc.add_paragraph("Methodology notes live on the second page.")
    doc.save(path)


def make_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Region", "Q1", "Q2", "Q3", "Total"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    rows = [("North", 100, 110, 120), ("South", 80, 90, 95), ("East", 60, 65, 70)]
    for i, (region, *values) in enumerate(rows, start=2):
        ws.append([region, *values, f"=SUM(B{i}:D{i})"])
    ws.append(["Total", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)", "=SUM(E2:E4)"])
    ws.merge_cells("A7:E7")
    ws["A7"] = "Figures in thousands. 单位：千元"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Assumption"
    notes["B1"] = "Value"
    notes["A2"] = "Growth"
    notes["B2"] = 0.12
    notes["B2"].number_format = "0%"
    wb.save(path)


def make_pptx(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Product Launch"
    s1.placeholders[1].text = "Plan for the autumn release"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Timeline"
    body = s2.placeholders[1].text_frame
    body.text = "September: beta"
    body.add_paragraph().text = "October: launch"
    body.add_paragraph().text = "十一月：复盘"
    box = s2.shapes.add_textbox(Inches(6), Inches(5.5), Inches(3), Inches(1))
    box.text_frame.text = "Owner: Dana"
    box.text_frame.paragraphs[0].runs[0].font.size = PptPt(14)
    s2.notes_slide.notes_text_frame.text = "Mention the beta feedback numbers."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Budget"
    rows, cols = 3, 3
    table = s3.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(8), Inches(1.5)).table
    for r, row in enumerate([("Item", "Cost", "Owner"), ("Ads", "40k", "Lee"), ("Events", "25k", "Kim")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    prs.save(path)


def make_pdf(path):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(path, pagesize=A4)
    c.setTitle("Research Paper")
    width, height = A4
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, height - 90, "Research Paper")
    c.setFont("Helvetica", 12)
    lines = [
        "Abstract: we study the effect of caching on latency.",
        "Results show a 40 percent reduction in p95 latency.",
    ]
    for i, line in enumerate(lines):
        c.drawString(72, height - 130 - i * 18, line)
    c.setFont("STSong-Light", 12)
    c.drawString(72, height - 180, "中文摘要：缓存显著降低延迟。")
    c.linkURL("https://example.com/paper", (72, height - 220, 300, height - 200))
    c.setFont("Helvetica", 12)
    c.drawString(72, height - 215, "https://example.com/paper")
    c.bookmarkPage("intro")
    c.addOutlineEntry("Introduction", "intro", level=0)
    c.showPage()
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 90, "Method")
    c.setFont("Helvetica", 12)
    c.drawString(72, height - 120, "We replay production traces against three cache sizes.")
    c.bookmarkPage("method")
    c.addOutlineEntry("Method", "method", level=0)
    c.showPage()
    c.save()


if __name__ == "__main__":
    make_docx("report.docx")
    make_xlsx("budget.xlsx")
    make_pptx("launch.pptx")
    make_pdf("paper.pdf")
